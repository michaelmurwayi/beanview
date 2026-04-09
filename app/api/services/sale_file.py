import os
import zipfile
from io import BytesIO
import logging
import shutil
from datetime import datetime

from django.conf import settings
from django.http import FileResponse
from rest_framework.response import Response
from rest_framework import status
from rest_framework.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from ..models import Coffee
from ..serializers import CoffeeSerializer

logger = logging.getLogger(__name__)


def generate_sales_file(request):
    TEMPLATE_PATH = os.path.join(settings.MEDIA_ROOT, "templates", "sale_summary_template.xlsx")
    START_ROW = 20

    try:
        sale_data = request.data.get("saleNumber")
        if not sale_data:
            raise ValidationError("saleNumber is required")

        sale_number = sale_data.get("sale number")
        if not sale_number:
            raise ValidationError("Sale Number must not be empty")

        logger.info(f"Generating sale file for sale: {sale_number}")

        base_dir = os.path.join(settings.MEDIA_ROOT, "summaries")
        os.makedirs(base_dir, exist_ok=True)

        generated_files = []

        # ✅ Fix here: select_related only on valid foreign keys
        coffees = Coffee.objects.filter(sale=sale_number).select_related(
            "farmer", "mill", "warehouse", "status"
        )

        serializer = CoffeeSerializer(coffees, many=True)
        coffees_data = serializer.data

        logger.info(f"Fetched {len(coffees_data)} coffee records")

        # Group coffees by farmer mark
        grouped = {}
        for coffee in coffees_data:
            mark = coffee.get("farmer", {}).get("mark")
            if not mark:
                logger.warning("Skipping coffee record with no farmer mark")
                continue
            grouped.setdefault(mark, []).append(coffee)

        logger.info(f"Grouped into {len(grouped)} marks")

        # Generate Excel files
        for mark, records in grouped.items():
            if not records:
                continue

            mark_dir = os.path.join(base_dir, mark.replace(" ", "_"))
            os.makedirs(mark_dir, exist_ok=True)

            filename = f"{mark.replace(' ', '_')}_summary_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            file_path = os.path.join(mark_dir, filename)

            shutil.copy(TEMPLATE_PATH, file_path)

            wb = load_workbook(file_path, data_only=False, keep_vba=True)
            ws = wb.active

            ws["B1"].value = records[0]["farmer"]["code"]
            ws["B2"].value = mark
            ws["B3"].value = sale_number

            for row_offset, record in enumerate(records, start=1):
                
                row = START_ROW + row_offset
                values = [
                    record.get("outturn"),
                    record.get("season"),
                    record.get("bags"),
                    record.get("pockets"),
                    record.get("weight"),
                    record.get("grade"),
                    record.get("price"),
                    record.get("gross_value"),
                    record.get("warehouse_charges"),
                    record.get("brokerage_charges"),
                    record.get("milling_charges"),
                    record.get("mill"),
                    record.get("export_charges"),
                    record.get("transport_charges"),
                    record.get("broker_transport"),
                    record.get("net_value"),
                    record.get("buyer")
                ]
                for col_index, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col_index)
                    if isinstance(cell, MergedCell):
                        continue
                    cell.value = value

            wb.save(file_path)
            generated_files.append(file_path)
            logger.info(f"Generated file: {file_path}")

        if not generated_files:
            raise ValidationError("No files were generated")

        # Create ZIP
        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, "w") as zip_file:
            for file_path in generated_files:
                zip_file.write(file_path, os.path.basename(file_path))
        zip_buffer.seek(0)

        logger.info("ZIP file generated successfully")

        return FileResponse(
            zip_buffer,
            as_attachment=True,
            filename=f"stock_summaries_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip",
            content_type="application/zip",
        )

    except ValidationError as e:
        logger.warning(str(e))
        return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    except Exception as e:
        logger.exception("Internal server error during sale file generation")
        return Response({"error": "Internal server error"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)